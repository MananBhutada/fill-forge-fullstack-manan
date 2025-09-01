import os, re, uuid, json
from fastapi import FastAPI, HTTPException, Depends, UploadFile, File, Form
from fastapi.security import OAuth2PasswordBearer, OAuth2PasswordRequestForm
from fastapi.middleware.cors import CORSMiddleware
from pydantic import BaseModel
from passlib.context import CryptContext
from datetime import datetime, timedelta
import jwt
from typing import List, Optional
from sqlalchemy import create_engine, Column, Integer, String, DateTime, ForeignKey, Enum, Text
from sqlalchemy.orm import sessionmaker, declarative_base, relationship
from sqlalchemy.exc import IntegrityError
from dotenv import load_dotenv
from pathlib import Path
from rapidfuzz import fuzz
from docx import Document

load_dotenv()
SECRET_KEY = os.environ.get("SECRET_KEY","changeme")
ALGORITHM = "HS256"
ACCESS_TOKEN_EXPIRE_MINUTES = int(os.environ.get("ACCESS_TOKEN_EXPIRE_MINUTES",60))

DATABASE_URL = os.environ.get("DATABASE_URL","sqlite:///./test.db")

# SQLAlchemy setup
engine = create_engine(DATABASE_URL, echo=False, future=True)
SessionLocal = sessionmaker(bind=engine, autoflush=False, autocommit=False)
Base = declarative_base()

class User(Base):
    __tablename__ = "users"
    id = Column(Integer, primary_key=True, index=True)
    email = Column(String(200), unique=True, index=True, nullable=False)
    password_hash = Column(String(200), nullable=False)
    name = Column(String(200), nullable=True)
    dob = Column(String(50), nullable=True)
    phone = Column(String(50), nullable=True)
    address = Column(Text, nullable=True)
    father_name = Column(String(200), nullable=True)
    mother_name = Column(String(200), nullable=True)
    created_at = Column(DateTime, default=datetime.utcnow)

class DocumentRecord(Base):
    __tablename__ = "documents"
    id = Column(Integer, primary_key=True, index=True)
    user_id = Column(Integer, ForeignKey("users.id", ondelete="CASCADE"))
    label = Column(String(50))
    file_path = Column(String(500))
    mime_type = Column(String(100))
    created_at = Column(DateTime, default=datetime.utcnow)

class FormRecord(Base):
    __tablename__ = "forms"
    id = Column(Integer, primary_key=True, index=True)
    user_id = Column(Integer, ForeignKey("users.id", ondelete="CASCADE"))
    original_filename = Column(String(255))
    stored_path = Column(String(500))
    status = Column(String(50), default="UPLOADED")
    created_at = Column(DateTime, default=datetime.utcnow)

class FormFill(Base):
    __tablename__ = "form_fills"
    id = Column(Integer, primary_key=True, index=True)
    form_id = Column(Integer, ForeignKey("forms.id"))
    mapping_json = Column(Text)
    output_path = Column(String(500))
    created_at = Column(DateTime, default=datetime.utcnow)

class Synonym(Base):
    __tablename__ = "synonyms"
    id = Column(Integer, primary_key=True, index=True)
    canonical_key = Column(String(200), index=True)
    variants_json = Column(Text, default="[]")

Base.metadata.create_all(bind=engine)

pwd_context = CryptContext(schemes=["bcrypt"], deprecated="auto")
oauth2_scheme = OAuth2PasswordBearer(tokenUrl="auth/login")

app = FastAPI(title="FillForge Auto - Backend")

app.add_middleware(
    CORSMiddleware,
    allow_origins=["http://localhost:3000"],
    allow_credentials=True,
    allow_methods=["*"],
    allow_headers=["*"],
)

UPLOAD_DIR = Path("./uploads")
FILLED_DIR = Path("./filled")
UPLOAD_DIR.mkdir(exist_ok=True)
FILLED_DIR.mkdir(exist_ok=True)

# ---------- Utility ----------
def get_db():
    db = SessionLocal()
    try:
        yield db
    finally:
        db.close()

def verify_password(plain_password, hashed):
    return pwd_context.verify(plain_password, hashed)

def get_password_hash(password):
    return pwd_context.hash(password)

def create_access_token(data: dict, expires_delta: Optional[timedelta] = None):
    to_encode = data.copy()
    if expires_delta:
        expire = datetime.utcnow() + expires_delta
    else:
        expire = datetime.utcnow() + timedelta(minutes=ACCESS_TOKEN_EXPIRE_MINUTES)
    to_encode.update({"exp": expire})
    encoded_jwt = jwt.encode(to_encode, SECRET_KEY, algorithm=ALGORITHM)
    return encoded_jwt

def decode_token(token: str):
    try:
        payload = jwt.decode(token, SECRET_KEY, algorithms=[ALGORITHM])
        return payload
    except jwt.PyJWTError:
        return None

def get_current_user(token: str = Depends(oauth2_scheme), db=Depends(get_db)):
    payload = decode_token(token)
    if not payload:
        raise HTTPException(status_code=401, detail="Invalid token")
    user = db.query(User).filter(User.id == payload.get("sub")).first()
    if not user:
        raise HTTPException(status_code=404, detail="User not found")
    return user

# ---------- Schemas ----------
class RegisterIn(BaseModel):
    email: str
    password: str
class TokenOut(BaseModel):
    access_token: str
    token_type: str = "bearer"
class ProfileOut(BaseModel):
    email: str
    name: Optional[str]
    dob: Optional[str]
    phone: Optional[str]
    address: Optional[str]
    father_name: Optional[str]
    mother_name: Optional[str]

# ---------- Auth ----------
@app.post("/auth/register", response_model=TokenOut)
def register(payload: RegisterIn, db=Depends(get_db)):
    hashed = get_password_hash(payload.password)
    user = User(email=payload.email, password_hash=hashed)
    db.add(user)
    try:
        db.commit()
        db.refresh(user)
    except IntegrityError:
        db.rollback()
        raise HTTPException(status_code=400, detail="Email already registered")
    token = create_access_token({"sub": user.id})
    return {"access_token": token, "token_type": "bearer"}

@app.post("/auth/login", response_model=TokenOut)
def login(form_data: OAuth2PasswordRequestForm = Depends(), db=Depends(get_db)):
    user = db.query(User).filter(User.email == form_data.username).first()
    if not user or not verify_password(form_data.password, user.password_hash):
        raise HTTPException(status_code=401, detail="Invalid credentials")
    token = create_access_token({"sub": user.id})
    return {"access_token": token, "token_type": "bearer"}

@app.get("/auth/me", response_model=ProfileOut)
def me(current=Depends(get_current_user)):
    return {
        "email": current.email,
        "name": current.name,
        "dob": current.dob,
        "phone": current.phone,
        "address": current.address,
        "father_name": current.father_name,
        "mother_name": current.mother_name
    }

# ---------- Profile ----------
@app.put("/profile", response_model=ProfileOut)
def update_profile(profile: ProfileOut, current=Depends(get_current_user), db=Depends(get_db)):
    u = db.query(User).filter(User.id == current.id).first()
    if not u:
        raise HTTPException(status_code=404, detail="User not found")
    u.name = profile.name
    u.dob = profile.dob
    u.phone = profile.phone
    u.address = profile.address
    u.father_name = profile.father_name
    u.mother_name = profile.mother_name
    db.commit()
    return profile

# ---------- Documents ----------
@app.post("/documents")
def upload_document(label: str = Form(...), file: UploadFile = File(...), current=Depends(get_current_user), db=Depends(get_db)):
    filename = f"{uuid.uuid4().hex}_{file.filename}"
    dest = UPLOAD_DIR / filename
    with open(dest, "wb") as f:
        f.write(file.file.read())
    rec = DocumentRecord(user_id=current.id, label=label, file_path=str(dest), mime_type=file.content_type)
    db.add(rec); db.commit(); db.refresh(rec)
    return {"id": rec.id, "label": rec.label, "file_path": rec.file_path}

@app.get("/documents")
def list_documents(current=Depends(get_current_user), db=Depends(get_db)):
    docs = db.query(DocumentRecord).filter(DocumentRecord.user_id == current.id).all()
    return [{"id":d.id,"label":d.label,"file_path":d.file_path} for d in docs]

@app.delete("/documents/{doc_id}")
def delete_document(doc_id: int, current=Depends(get_current_user), db=Depends(get_db)):
    d = db.query(DocumentRecord).filter(DocumentRecord.id==doc_id, DocumentRecord.user_id==current.id).first()
    if not d:
        raise HTTPException(status_code=404, detail="Document not found")
    try:
        os.remove(d.file_path)
    except Exception:
        pass
    db.delete(d); db.commit()
    return {"ok": True}

# ---------- Utils: docx placeholders ----------
PLACEHOLDER_RE = re.compile(r"\{\{([A-Z0-9_]+)\}\}")

def extract_placeholders(doc_path: str):
    doc = Document(doc_path)
    keys = set()
    for p in doc.paragraphs:
        for run in p.runs:
            m = PLACEHOLDER_RE.findall(run.text)
            for k in m: keys.add(k)
    # also in tables
    for table in doc.tables:
        for row in table.rows:
            for cell in row.cells:
                for p in cell.paragraphs:
                    for run in p.runs:
                        m = PLACEHOLDER_RE.findall(run.text)
                        for k in m: keys.add(k)
    return sorted(list(keys))

def replace_placeholders(doc_path: str, replacements: dict, out_path: str):
    doc = Document(doc_path)
    for p in doc.paragraphs:
        for run in p.runs:
            text = run.text
            for k,v in replacements.items():
                text = text.replace(f"{{{{{k}}}}}", str(v) if v is not None else "")
            run.text = text
    for table in doc.tables:
        for row in table.rows:
            for cell in row.cells:
                for p in cell.paragraphs:
                    for run in p.runs:
                        text = run.text
                        for k,v in replacements.items():
                            text = text.replace(f"{{{{{k}}}}}", str(v) if v is not None else "")
                        run.text = text
    doc.save(out_path)

# ---------- Form endpoints ----------
@app.post("/forms/upload")
def upload_form(file: UploadFile = File(...), current=Depends(get_current_user), db=Depends(get_db)):
    if not file.filename.lower().endswith(".docx"):
        raise HTTPException(status_code=400, detail="Only .docx allowed for now")
    filename = f"{uuid.uuid4().hex}_{file.filename}"
    dest = UPLOAD_DIR / filename
    with open(dest, "wb") as f:
        f.write(file.file.read())
    rec = FormRecord(user_id=current.id, original_filename=file.filename, stored_path=str(dest))
    db.add(rec); db.commit(); db.refresh(rec)
    placeholders = extract_placeholders(str(dest))
    return {"form_id": rec.id, "placeholders": placeholders}

@app.post("/forms/fill")
def fill_form(form_id: int = Form(...), overrides: Optional[str] = Form(None), current=Depends(get_current_user), db=Depends(get_db)):
    rec = db.query(FormRecord).filter(FormRecord.id==form_id, FormRecord.user_id==current.id).first()
    if not rec:
        raise HTTPException(status_code=404, detail="Form not found")
    placeholders = extract_placeholders(rec.stored_path)
    # build replacements from user profile and synonyms and fuzzy match
    replacements = {}
    user = current
    # canonical map from user model
    canonical = {
        "NAME": user.name or "",
        "DOB": user.dob or "",
        "PHONE": user.phone or "",
        "ADDRESS": user.address or "",
        "FATHER_NAME": user.father_name or "",
        "MOTHER_NAME": user.mother_name or "",
        "EMAIL": user.email or ""
    }
    # load synonyms table
    syns = db.query(Synonym).all()
    synmap = {}
    for s in syns:
        try:
            variants = json.loads(s.variants_json)
        except Exception:
            variants = []
        synmap[s.canonical_key] = variants
    # helper to find best match
    def find_value_for(key):
        # exact
        if key in canonical: return canonical[key]
        # synonyms
        for can, variants in synmap.items():
            if key == can: return canonical.get(can,"")
            for v in variants:
                if key.upper() == v.upper() or fuzz.ratio(key, v) >= 85:
                    return canonical.get(can, "")
        # fuzzy against canonical labels
        best = None; best_score = 0
        for can in canonical.keys():
            score = fuzz.ratio(key, can)
            if score > best_score:
                best_score = score; best = can
        if best_score >= 85:
            return canonical.get(best,"")
        return ""
    for k in placeholders:
        replacements[k] = find_value_for(k)
    # apply overrides if provided (overrides expected as JSON string mapping)
    if overrides:
        try:
            ov = json.loads(overrides)
            replacements.update(ov)
        except Exception:
            pass
    out_name = f"{uuid.uuid4().hex}_filled.docx"
    out_path = FILLED_DIR / out_name
    replace_placeholders(rec.stored_path, replacements, str(out_path))
    ff = FormFill(form_id=rec.id, mapping_json=json.dumps(replacements), output_path=str(out_path))
    rec.status = "FILLED"
    db.add(ff); db.commit()
    return {"filled_docx": str(out_path), "mapping": replacements}

@app.get("/health")
def health():
    return {"status":"ok","time": datetime.utcnow().isoformat()}